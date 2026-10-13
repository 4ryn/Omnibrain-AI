import csv
import io
from pathlib import Path

from app.core.errors import ExtractionError

SUPPORTED_EXTENSIONS = {".pdf", ".txt", ".md", ".markdown", ".docx", ".csv"}
MAX_PDF_PAGES = 500
MAX_CSV_ROWS = 20000
MAX_DOCX_UNCOMPRESSED = 100 * 1024 * 1024


def _decode(data: bytes) -> str:
    for enc in ("utf-8-sig", "utf-16"):
        try:
            return data.decode(enc)
        except UnicodeError:
            continue
    return data.decode("latin-1")


def _pdf(data: bytes) -> str:
    from pypdf import PdfReader
    from pypdf.errors import PyPdfError
    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted and not reader.decrypt(""):
            raise ExtractionError("This PDF is password-protected. Upload an unprotected copy.")
        parts = [(page.extract_text() or "") for page in reader.pages[:MAX_PDF_PAGES]]
    except ExtractionError:
        raise
    except (PyPdfError, ValueError, KeyError, OSError) as exc:
        raise ExtractionError("This PDF could not be read. It may be corrupted.") from exc
    text = "\n\n".join(parts).strip()
    if not text:
        raise ExtractionError("No text found in this PDF. Scanned/image-only PDFs are not supported yet.")
    return text


def _docx(data: bytes) -> str:
    import docx
    import zipfile
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as zf:
            if sum(i.file_size for i in zf.infolist()) > MAX_DOCX_UNCOMPRESSED:
                raise ExtractionError("This Word document expands to an unreasonable size and was rejected.")
    except zipfile.BadZipFile as exc:
        raise ExtractionError("This Word document could not be read. Only .docx is supported.") from exc
    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:
        raise ExtractionError("This Word document could not be read. Only .docx is supported.") from exc
    parts = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n\n".join(parts)


def _csv(data: bytes) -> str:
    text = _decode(data)
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    rows = csv.reader(io.StringIO(text), dialect)
    try:
        header = next(rows)
    except StopIteration:
        return ""
    lines = []
    for n, row in enumerate(rows):
        if n >= MAX_CSV_ROWS:
            break
        pairs = [f"{h.strip()}: {v.strip()}" for h, v in zip(header, row) if v.strip()]
        if pairs:
            lines.append("; ".join(pairs))
    return "\n\n".join(lines)


def extract_text(filename: str, data: bytes) -> str:
    ext = Path(filename).suffix.lower()
    if ext == ".pdf":
        return _pdf(data)
    if ext == ".docx":
        return _docx(data)
    if ext == ".csv":
        return _csv(data)
    if ext in {".txt", ".md", ".markdown"}:
        return _decode(data)
    raise ExtractionError(f"Unsupported file type '{ext}'. Supported: PDF, TXT, MD, DOCX, CSV.")
