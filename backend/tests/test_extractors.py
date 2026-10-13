import io

import docx
import pytest

from app.core.errors import ExtractionError
from app.ingestion.extractors import extract_text


def _minimal_pdf(text: str) -> bytes:
    objs = [
        "<< /Type /Catalog /Pages 2 0 R >>",
        "<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
        "/Resources << /Font << /F1 5 0 R >> >> >>",
        None,
        "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET"
    objs[3] = f"<< /Length {len(stream)} >>\nstream\n{stream}\nendstream"
    out, offsets = b"%PDF-1.4\n", []
    for i, o in enumerate(objs, 1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n{o}\nendobj\n".encode()
    xref = len(out)
    out += f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode()
    out += "".join(f"{o:010d} 00000 n \n" for o in offsets).encode()
    out += f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF".encode()
    return out


def test_pdf_text_extraction():
    assert "Our services include consulting" in extract_text("a.pdf", _minimal_pdf("Our services include consulting"))


def test_corrupt_pdf_raises_user_facing_error():
    with pytest.raises(ExtractionError):
        extract_text("a.pdf", b"%PDF-1.4 garbage")


def test_txt_and_markdown():
    assert extract_text("a.txt", "héllo".encode()) == "héllo"
    assert "# Title" in extract_text("a.md", b"# Title\nbody")


def test_docx_paragraphs_and_tables():
    d = docx.Document()
    d.add_paragraph("We offer web design.")
    t = d.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text, t.rows[0].cells[1].text = "Plan", "Pro"
    buf = io.BytesIO()
    d.save(buf)
    text = extract_text("a.docx", buf.getvalue())
    assert "web design" in text and "Plan | Pro" in text


def test_csv_rows_become_labelled_lines():
    text = extract_text("a.csv", b"name,price\nBasic,10\nPro,20\n")
    assert "name: Basic; price: 10" in text and "name: Pro; price: 20" in text


def test_unsupported_type():
    with pytest.raises(ExtractionError):
        extract_text("a.exe", b"MZ")
